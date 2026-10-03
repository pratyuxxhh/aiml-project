from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor

from app.errors import AppError
from app.models.document import DocumentSection, DocumentSpecification
from app.models.formatting_spec import DocumentFormatting, TextStyle, TableSpec, ColumnSpec, RowStyleRule, Color
import re
from typing import Optional


def write_docx(
    spec: DocumentSpecification, output_path: Path, formatting: Optional[DocumentFormatting] = None
) -> Path:
    try:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        document = Document()
        _apply_page_setup(document, spec, formatting)
        _set_run_font_defaults(document, spec.font_size, formatting)

        for section in spec.sections:
            _render_section(document, section, spec.font_size, formatting)

        _add_footer(document, spec.title)
        document.save(str(output_path))
        return output_path
    except AppError:
        raise
    except Exception as exc:
        raise AppError(f"Failed to write DOCX: {exc}", status_code=500) from exc


def _apply_page_setup(document: Document, spec: DocumentSpecification, formatting: Optional[DocumentFormatting]) -> None:
    section = document.sections[0]
    left = 0.9
    right = 0.9
    top = 0.9
    bottom = 0.9
    if formatting and formatting.margins:
        if formatting.margins.left_mm is not None:
            left = formatting.margins.left_mm / 25.4
        if formatting.margins.right_mm is not None:
            right = formatting.margins.right_mm / 25.4
        if formatting.margins.top_mm is not None:
            top = formatting.margins.top_mm / 25.4
        if formatting.margins.bottom_mm is not None:
            bottom = formatting.margins.bottom_mm / 25.4
    section.left_margin = Inches(left)
    section.right_margin = Inches(right)
    section.top_margin = Inches(top)
    section.bottom_margin = Inches(bottom)
    orient = getattr(spec, "orientation", "portrait")
    if formatting and formatting.orientation:
        orient = formatting.orientation
    if orient == "landscape":
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width


def _set_run_font_defaults(document: Document, font_size: int, formatting: Optional[DocumentFormatting]) -> None:
    style = document.styles["Normal"]
    if formatting and formatting.body:
        body = formatting.body
        if body.font_family:
            style.font.name = body.font_family
        if body.size_pt:
            style.font.size = Pt(body.size_pt)
        else:
            style.font.size = Pt(font_size)
        if body.color and body.color.hex:
            rgb = _hex_to_rgb(body.color.hex)
            if rgb:
                style.font.color.rgb = RGBColor(*rgb)
    else:
        style.font.name = "Calibri"
        style.font.size = Pt(font_size)
        style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)


def _render_section(document: Document, section: DocumentSection, font_size: int, formatting: Optional[DocumentFormatting]) -> None:
    if section.type == "heading":
        paragraph = document.add_heading(section.text or "", level=min(max(section.level, 1), 3))
        if formatting and formatting.headings:
            for run in paragraph.runs:
                _apply_textstyle_to_run(run, formatting.headings)
        else:
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
        return
    if section.type == "paragraph":
        paragraph = document.add_paragraph(section.text or "")
        paragraph.paragraph_format.space_after = Pt(10)
        if formatting and formatting.body and formatting.body.alignment:
            align = formatting.body.alignment
            if align == "center":
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif align == "right":
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif align == "justify":
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if formatting and formatting.body:
            for run in paragraph.runs:
                _apply_textstyle_to_run(run, formatting.body)
        return
    if section.type == "summary":
        if not section.items:
            return
        table = document.add_table(rows=1, cols=len(section.items))
        _shade_header(table.rows[0], formatting)
        for i, item in enumerate(section.items):
            cell = table.rows[0].cells[i]
            cell.text = f"{item.label}: {item.value}"
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.bold = True
                    run.font.size = Pt(font_size)
        document.add_paragraph("")
        return
    if section.type == "table":
        _add_data_table(document, section, font_size, formatting)
        return
    if section.type == "cards":
        for card in section.cards:
            heading = document.add_heading(card.title, level=3)
            for run in heading.runs:
                run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
            if card.fields:
                table = document.add_table(rows=len(card.fields), cols=2)
                table.style = "Table Grid"
                for i, field in enumerate(card.fields):
                    table.rows[i].cells[0].text = field.label
                    table.rows[i].cells[1].text = field.value
                    for run in table.rows[i].cells[0].paragraphs[0].runs:
                        run.bold = True
            document.add_paragraph("")


def _add_data_table(document: Document, section: DocumentSection, font_size: int, formatting: Optional[DocumentFormatting]) -> None:
    cols = max(len(section.headers), 1)
    table = document.add_table(rows=1 + len(section.rows), cols=cols)
    table.style = "Table Grid"
    table.autofit = False
    _set_table_fixed_layout(table)
    header_row = table.rows[0]
    _shade_header(header_row, formatting)
    for i, header in enumerate(section.headers):
        cell = header_row.cells[i]
        cell.text = header
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(font_size)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for r, row in enumerate(section.rows, start=1):
        for c in range(cols):
            value = row[c] if c < len(row) else ""
            table.rows[r].cells[c].text = value
            for paragraph in table.rows[r].cells[c].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(max(font_size - 1, 8))

    available_width = int(
        document.sections[0].page_width
        - document.sections[0].left_margin
        - document.sections[0].right_margin
    )
    _set_table_width(table, available_width)
    if section.column_widths and len(section.column_widths) == len(table.columns):
        total_width = sum(section.column_widths) or 100
        for index, width_percent in enumerate(section.column_widths):
            width = max(int(available_width * width_percent / total_width), 1)
            _set_column_width(table, index, width)
    else:
        even = max(available_width // cols, 1)
        for index in range(cols):
            _set_column_width(table, index, even)

    if formatting and formatting.sections:
        for sec in formatting.sections:
            if sec.type == "table" and sec.id:
                content = sec.content
                if hasattr(content, "table_spec") and content.table_spec:
                    t_spec: TableSpec = content.table_spec
                    if t_spec.table_alignment:
                        if t_spec.table_alignment == "center":
                            table.alignment = WD_TABLE_ALIGNMENT.CENTER
                        elif t_spec.table_alignment == "right":
                            table.alignment = WD_TABLE_ALIGNMENT.RIGHT
                        else:
                            table.alignment = WD_TABLE_ALIGNMENT.LEFT
                    if t_spec.header_style:
                        for cell in header_row.cells:
                            for paragraph in cell.paragraphs:
                                for run in paragraph.runs:
                                    _apply_textstyle_to_run(run, t_spec.header_style)
                    if t_spec.columns:
                        for i, col_spec in enumerate(t_spec.columns):
                            if i >= len(table.columns):
                                break
                            if col_spec.width_mm:
                                inches = col_spec.width_mm / 25.4
                                table.columns[i].width = Inches(inches)
                            elif col_spec.width_percent:
                                available_width = document.sections[0].page_width - document.sections[0].left_margin - document.sections[0].right_margin
                                width = int(available_width * col_spec.width_percent / 100)
                                _set_column_width(table, i, width)
                    if t_spec.row_style_rules:
                        for rule in t_spec.row_style_rules:
                            _apply_row_style_rule(table, rule, section.headers)


def _shade_header(row, formatting: Optional[DocumentFormatting] = None) -> None:
    color_hex = None
    if formatting and formatting.sections:
        for sec in formatting.sections:
            if sec.type == "table" and getattr(sec.content, "table_spec", None):
                tspec = sec.content.table_spec
                if tspec and tspec.header_style and tspec.header_style.background_color and tspec.header_style.background_color.hex:
                    color_hex = tspec.header_style.background_color.hex
                    break
    if color_hex:
        rgb = _hex_to_rgb(color_hex)
        fill = rgb and "{:02X}{:02X}{:02X}".format(*rgb) or "1A365D"
    else:
        fill = "1A365D"
    for cell in row.cells:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), fill)
        shading.set(qn("w:val"), "clear")
        cell._tc.get_or_add_tcPr().append(shading)


def _set_table_fixed_layout(table) -> None:
    tbl_pr = table._tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        table._tbl.insert(0, tbl_pr)
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def _set_table_width(table, width: int) -> None:
    tbl_pr = table._tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        table._tbl.insert(0, tbl_pr)
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    twips = max(int(width / 635), 1)
    tbl_w.set(qn("w:w"), str(twips))
    tbl_w.set(qn("w:type"), "dxa")


def _set_column_width(table, index: int, width: int) -> None:
    column = table.columns[index]
    twips = max(int(width / 635), 1)
    column.width = width
    column._gridCol.set(qn("w:w"), str(twips))
    for cell in column.cells:
        cell.width = width
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_width = tc_pr.first_child_found_in("w:tcW")
        if tc_width is None:
            tc_width = OxmlElement("w:tcW")
            tc_pr.append(tc_width)
        tc_width.set(qn("w:w"), str(twips))
        tc_width.set(qn("w:type"), "dxa")
        no_wrap = tc_pr.first_child_found_in("w:noWrap")
        if no_wrap is not None:
            tc_pr.remove(no_wrap)


def _hex_to_rgb(hex_str: str):
    if not hex_str:
        return None
    s = hex_str.strip().lstrip("#")
    if len(s) == 8:
        s = s[2:]
    if len(s) != 6:
        return None
    try:
        return tuple(int(s[i:i + 2], 16) for i in (0, 2, 4))
    except ValueError:
        return None


def _apply_textstyle_to_run(run, style: TextStyle) -> None:
    if style.font_family:
        try:
            run.font.name = style.font_family
        except Exception:
            pass
    if style.size_pt:
        try:
            run.font.size = Pt(style.size_pt)
        except Exception:
            pass
    if style.color and style.color.hex:
        rgb = _hex_to_rgb(style.color.hex)
        if rgb:
            run.font.color.rgb = RGBColor(*rgb)
    if style.bold is not None:
        run.bold = bool(style.bold)
    if style.italic is not None:
        run.italic = bool(style.italic)
    if style.underline is not None:
        run.underline = bool(style.underline)


def _apply_row_style_rule(table, rule: RowStyleRule, headers: list[str]) -> None:
    try:
        col_idx = headers.index(rule.condition.column)
    except ValueError:
        return
    op = rule.condition.operator
    val = rule.condition.value
    for r in range(1, len(table.rows)):
        cell = table.rows[r].cells[col_idx]
        text = "".join(p.text for p in cell.paragraphs)
        try:
            lhs = float(text)
            rhs = float(val) if isinstance(val, (int, float, str)) and str(val).replace("%", "") != "" else val
        except Exception:
            lhs = text
            rhs = val
        matched = False
        if op == ">":
            try:
                matched = float(lhs) > float(rhs)
            except Exception:
                matched = False
        elif op == "<":
            try:
                matched = float(lhs) < float(rhs)
            except Exception:
                matched = False
        elif op == ">=":
            try:
                matched = float(lhs) >= float(rhs)
            except Exception:
                matched = False
        elif op == "<=":
            try:
                matched = float(lhs) <= float(rhs)
            except Exception:
                matched = False
        elif op == "==":
            matched = str(lhs) == str(rhs)
        elif op == "!=":
            matched = str(lhs) != str(rhs)
        elif op == "contains":
            matched = str(rhs) in str(lhs)
        elif op == "regex":
            try:
                matched = bool(re.search(str(rhs), str(lhs)))
            except Exception:
                matched = False
        if matched:
            action = rule.action
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    if action.font_color and action.font_color.hex:
                        rgb = _hex_to_rgb(action.font_color.hex)
                        if rgb:
                            run.font.color.rgb = RGBColor(*rgb)
                    if action.bold:
                        run.bold = True
                    if action.italic:
                        run.italic = True
                    if action.underline:
                        run.underline = True
            if action.highlight_bg and action.highlight_bg.hex:
                fill = _hex_to_rgb(action.highlight_bg.hex)
                if fill:
                    fill_str = "{:02X}{:02X}{:02X}".format(*fill)
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:fill"), fill_str)
                    shading.set(qn("w:val"), "clear")
                    cell._tc.get_or_add_tcPr().append(shading)


def _add_footer(document: Document, title: str) -> None:
    footer = document.sections[0].footer
    paragraph = footer.paragraphs[0]
    paragraph.text = title
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run("  |  Generated by Excel → Word")
    run.italic = True
