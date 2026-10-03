from pathlib import Path

from docx import Document

from app.models.document import DocumentEdits
from app.services.excel_service import extract_dataset
from app.services.template_service import apply_analysis, build_specification
from app.services.classification_service import classify_columns
from app.services.docx_service import write_docx
from app.services.preview_service import render_preview_html

SAMPLE = Path(__file__).resolve().parents[1] / "sample" / "placement.csv"


def _dataset():
    dataset, _ = extract_dataset(SAMPLE, source_file="placement.csv", sheet=None)
    analysis = classify_columns([c.original_name for c in dataset.columns])
    return apply_analysis(dataset, analysis)


def test_all_templates_build_spec_with_resolved_content():
    dataset = _dataset()
    template_ids = (
        "generic_table",
        "professional_report",
        "profile_cards",
        "placement_summary",
        "academic_marksheet",
        "attendance_register",
    )
    for template_id in template_ids:
        spec = build_specification(dataset, template_id, DocumentEdits(show_summary=True))
        assert spec.template == template_id
        types = {s.type for s in spec.sections}
        assert "heading" in types
        if template_id == "profile_cards":
            assert "cards" in types
            cards = next(s for s in spec.sections if s.type == "cards")
            assert len(cards.cards) == 3
        else:
            assert "table" in types
            table = next(s for s in spec.sections if s.type == "table")
            assert table.headers[0] == "Name"
            assert table.rows[0][0] == "Pratyush"


def test_edits_hide_summary_and_columns():
    dataset = _dataset()
    spec = build_specification(
        dataset,
        "generic_table",
        DocumentEdits(
            title="Custom Title",
            show_summary=False,
            selected_columns=["name", "usn"],
        ),
    )
    assert spec.title == "Custom Title"
    assert all(s.type != "summary" for s in spec.sections)
    table = next(s for s in spec.sections if s.type == "table")
    assert table.headers == ["Name", "USN"]


def test_preview_and_docx_use_same_spec(tmp_path: Path):
    dataset = _dataset()
    spec = build_specification(dataset, "professional_report", DocumentEdits(title="2026 Placement Report"))
    html = render_preview_html(spec)
    assert "2026 Placement Report" in html
    assert "Pratyush" in html
    out = tmp_path / "out.docx"
    write_docx(spec, out)
    assert out.exists()
    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs]
    assert any("2026 Placement Report" in t for t in texts)
    table_texts = []
    for table in doc.tables:
        for row in table.rows:
            table_texts.extend(cell.text for cell in row.cells)
    assert "Pratyush" in table_texts


def test_column_widths_are_rendered_in_preview_and_docx(tmp_path: Path):
    dataset = _dataset()
    widths = [10, 15, 20, 25, 30]
    spec = build_specification(
        dataset,
        "generic_table",
        DocumentEdits(column_widths=widths, selected_columns=["name", "usn", "cgpa", "package", "company"]),
    )
    table = next(section for section in spec.sections if section.type == "table")
    assert table.column_widths == widths
    html = render_preview_html(spec)
    assert 'style="width: 10.0%;"' in html
    assert 'data-col-width="10.0"' in html

    out = tmp_path / "widths.docx"
    write_docx(spec, out)
    doc = Document(str(out))
    exported_table = next(
        table for table in doc.tables
        if any(cell.text == "Pratyush" for row in table.rows for cell in row.cells)
    )
    exported_widths = [column.width for column in exported_table.columns]
    assert exported_widths[0] < exported_widths[1] < exported_widths[2]
    assert exported_widths[2] < exported_widths[3] < exported_widths[4]


def test_mismatched_column_widths_are_ignored():
    dataset = _dataset()
    spec = build_specification(
        dataset,
        "generic_table",
        DocumentEdits(column_widths=[20, 30, 50], selected_columns=["name", "usn"]),
    )
    table = next(section for section in spec.sections if section.type == "table")
    assert table.headers == ["Name", "USN"]
    assert table.column_widths == []


def test_profile_cards_export_as_three_column_grid(tmp_path: Path):
    dataset = _dataset()
    spec = build_specification(dataset, "profile_cards", DocumentEdits())
    out = tmp_path / "cards.docx"

    write_docx(spec, out)

    doc = Document(str(out))
    card_grid = next(
        table for table in doc.tables
        if any(cell.text.startswith("Pratyush") for row in table.rows for cell in row.cells)
    )
    assert len(card_grid.columns) == 3
    assert len(card_grid.rows) == 1
    assert "Pratyush" in card_grid.cell(0, 0).text
    assert "Ananya" in card_grid.cell(0, 1).text
    assert "Rohan" in card_grid.cell(0, 2).text
